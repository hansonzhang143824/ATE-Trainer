# -*- coding: utf-8 -*-
"""Probe: list sheets/dimensions/headers of DFT workbooks via python (DLP-authorized plaintext read)."""
import sys, os, json

ROOT = r"D:/Newtest/DSH/ATE-Coding-Plat"
def p(*a):
    return os.path.join(ROOT, *a)

def read_bytes(path):
    with open(path, 'rb') as f:
        return f.read()

def decode(raw):
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'latin-1'):
        try:
            return raw.decode(enc), enc
        except Exception:
            continue
    return raw.decode('utf-8', errors='replace'), 'replace'

out = {}

# 1) DFT.csv sniff
csv_path = p('project', 'DALI', 'input', 'DFT.csv')
raw = read_bytes(csv_path)
txt, enc = decode(raw)
print('=== DFT.csv ===', len(raw), 'bytes, decoded as', enc)
print('--- first 3 lines ---')
for line in txt.splitlines()[:3]:
    print(repr(line[:400]))
print('--- TM600/TM601 matching lines ---')
for i, line in enumerate(txt.splitlines(), 1):
    if 'TM600' in line or 'TM601' in line or 'RDSON' in line.upper():
        print(i, repr(line[:400]))

# 2) xlsx sheets
try:
    import openpyxl
    for name in ('Dali_testmode.xlsx',):
        path = p('project', 'DALI', name)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        print('===', name, 'sheets:', wb.sheetnames)
        for sn in wb.sheetnames:
            ws = wb[sn]
            print(' sheet', sn, 'dims', ws.max_row, 'x', ws.max_column)
            if sn.lower().startswith('overview'):
                for r in ws.iter_rows(min_row=1, max_row=3, values_only=True):
                    print('   hdr:', [str(c)[:40] if c is not None else None for c in r])
        wb.close()
except Exception as e:
    print('xlsx error', type(e).__name__, e)
