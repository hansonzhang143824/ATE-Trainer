#!/usr/bin/env python3
"""Canonical project material locations for the ATE PTC pipeline.

Fixed layout inside the project directory:
  Input_GlobalMaterial/    read-only source material: DFT workbook, schematic CSV, CBIT table
  Output_Global_Material/  the only place DFT and schematic parsing products are written
  ErrorLog/                the only place role error logs are written

The DFT authority is the workbook in Input_GlobalMaterial whose filename contains
"DFT" or "testmode"; only its OVERVIEW sheet is read.
"""
from __future__ import annotations
import json
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

ROOT = Path(__file__).resolve().parents[1]
PROJECT_DIR = ROOT / 'project' / 'DALI'
INPUT_ROOT = PROJECT_DIR / 'Input_GlobalMaterial'
OUTPUT_ROOT = PROJECT_DIR / 'Output_Global_Material'
ERROR_ROOT = PROJECT_DIR / 'ErrorLog'

OVERVIEW_SHEET = 'OVERVIEW'
WORKBOOK_SUFFIXES = ('.xlsx', '.xlsm', '.xls')
SCHEMATIC_SUFFIXES = ('.csv',)
DFT_TOKENS = ('dft', 'testmode')
SCHEMATIC_TOKENS = ('sch',)
CBIT_TOKENS = ('cbit',)


def material_root(kind: str) -> Path:
    try:
        return {'input': INPUT_ROOT, 'output': OUTPUT_ROOT, 'error': ERROR_ROOT}[kind]
    except KeyError:
        raise ValueError('unknown material root %r' % (kind,))


def _candidates(root: Path, tokens, suffixes) -> list:
    found = []
    if not root.is_dir():
        return found
    for path in sorted(root.rglob('*')):
        if not path.is_file() or path.name.startswith('~$'):
            continue
        if suffixes and path.suffix.lower() not in suffixes:
            continue
        name = path.name.lower()
        if any(token in name for token in tokens):
            found.append(path)
    return found


def _inside(path: Path, root: Path) -> bool:
    resolved, base = Path(path).resolve(), Path(root).resolve()
    return resolved == base or base in resolved.parents


def ensure_within(path: Path, root: Path) -> Path:
    """Every material read/write must stay inside its declared root."""
    resolved, base = Path(path).resolve(), Path(root).resolve()
    if resolved != base and base not in resolved.parents:
        raise ValueError('%s is outside %s' % (resolved, base))
    return resolved


def _single(found: list, what: str, root: Path) -> Path:
    if len(found) != 1:
        raise ValueError('expected exactly one %s in %s; found %d: %s'
                         % (what, root, len(found), [p.name for p in found]))
    return found[0]


def workbook_candidates(root: Path | None = None) -> list:
    return _candidates(root or INPUT_ROOT, DFT_TOKENS, WORKBOOK_SUFFIXES)


def schematic_candidates(root: Path | None = None) -> list:
    return _candidates(root or INPUT_ROOT, SCHEMATIC_TOKENS, SCHEMATIC_SUFFIXES)


def cbit_candidates(root: Path | None = None) -> list:
    return _candidates(root or INPUT_ROOT, CBIT_TOKENS, WORKBOOK_SUFFIXES)


def discover_workbook(root: Path | None = None) -> Path:
    """The single DFT workbook in Input_GlobalMaterial."""
    root = (root or INPUT_ROOT).resolve()
    try:
        import project_info
        declared = (project_info.load() or {}).get('inputs', {}).get('dft')
        if isinstance(declared, str) and declared:
            candidate = project_info.resolve(declared)
            if candidate.is_file() and _inside(candidate, root):
                return candidate
    except Exception:
        pass
    return _single(workbook_candidates(root), 'DFT workbook (filename contains "DFT" or "testmode")', root)


def discover_schematic(root: Path | None = None) -> Path:
    """The single schematic source file in Input_GlobalMaterial."""
    root = (root or INPUT_ROOT).resolve()
    try:
        import project_info
        declared = (project_info.load() or {}).get('inputs', {}).get('schematic')
        if isinstance(declared, str) and declared:
            candidate = project_info.resolve(declared)
            if candidate.is_file() and _inside(candidate, root):
                return candidate
    except Exception:
        pass
    return _single(schematic_candidates(root), 'schematic source (filename contains "sch")', root)


def dft_output_dir(tm: str) -> Path:
    return OUTPUT_ROOT / 'dft' / str(tm).upper()


def schematic_output_dir() -> Path:
    return OUTPUT_ROOT / 'schematic'


def error_log_path(role: str, tm: str = '') -> Path:
    stem = str(role) + (('-' + str(tm).upper()) if tm else '')
    return ERROR_ROOT / (stem + '.log')


def overview_rows(workbook: Path) -> dict:
    """Read only the OVERVIEW sheet -> {Item: {column: value}}, None replaced by an empty string."""
    if openpyxl is None:
        raise ValueError('openpyxl is required to read the DFT workbook')
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        if OVERVIEW_SHEET not in wb.sheetnames:
            raise ValueError('workbook has no %s sheet' % OVERVIEW_SHEET)
        rows = list(wb[OVERVIEW_SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return {}
    header = [str(h).strip() if h is not None else '' for h in rows[0]]
    records = {}
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        item = str(row[0]).strip()
        if not item.startswith('TM'):
            continue
        record = {}
        for index, column in enumerate(header):
            value = row[index] if index < len(row) else None
            record[column] = '' if value is None else value
        records[item] = record
    return records


def rows_for(workbook: Path, tm: str) -> list:
    """All matching OVERVIEW rows, retaining duplicates for the one-row gate."""
    if openpyxl is None:
        raise ValueError('openpyxl is required to read the DFT workbook')
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        if OVERVIEW_SHEET not in wb.sheetnames:
            raise ValueError('workbook has no %s sheet' % OVERVIEW_SHEET)
        rows = wb[OVERVIEW_SHEET].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else '' for h in next(rows, ())]
        matches = []
        for row in rows:
            if not row or str(row[0]).strip().upper() != str(tm).upper():
                continue
            matches.append({column: '' if (index >= len(row) or row[index] is None) else row[index]
                            for index, column in enumerate(header)})
        return matches
    finally:
        wb.close()
